import json
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from chem_pdf_extractor.biomedical import extract_article, text_chunks
from chem_pdf_extractor.biology import ANTI_AGING_FIELDS, BIOLOGY_RULES
from chem_pdf_extractor.config import RuntimeDeps
from chem_pdf_extractor.enrichment import enrich_compound, enrich_rows
from chem_pdf_extractor.evaluation import evaluate
from chem_pdf_extractor.literature import build_pubmed_query, search_related_papers
from chem_pdf_extractor.science_skills import ScienceSkills, SkillError
from chem_pdf_extractor.sources import load_paper
from chem_pdf_extractor.validation import source_span, validate_experiment

SOURCE = '# Page 2\nQuercetin was tested in fibroblasts at 2 uM for 48 h using SA-beta-gal. No change was observed.'


def experiment(**overrides):
    row = dict(zip(("compound_name", "experiment_model", "treatment_dose", "treatment_duration",
                    "assay_method", "observed_result"),
                   ("Quercetin", "fibroblasts", "2 uM", "48 h", "SA-beta-gal", "No change was observed.")))
    row["field_evidence"] = json.dumps({k: SOURCE.split('\n', 1)[1] for k in row})
    row.update(overrides)
    return row


class FakeSkills:
    def __init__(self, cids=None, fail="", wrong_key=False):
        self.cids = [1] if cids is None else cids
        self.fail = fail
        self.wrong_key = wrong_key
        self.calls = []

    def call(self, db, *args):
        self.calls.append((db, args))
        if db == self.fail:
            raise SkillError('network error')
        command = args[0]
        if db == 'pubchem':
            if command == 'resolve':
                return {'identifiers': {'IdentifierList': {'CID': self.cids}},
                        'properties': {'PropertyTable': {'Properties': [{'CID': 1, 'InChIKey': 'FULL-KEY-N'}]}}}
            if command == 'properties':
                return {'PropertyTable': {'Properties': [{'CID': 1, 'MolecularWeight': '302'}]}}
            if command == 'synonyms':
                return {'InformationList': {'Information': [{'CID': 1, 'Synonym': ['Q']} ]}}
        if db == 'chembl':
            if command == 'molecule':
                return {'molecules': [{'molecule_chembl_id': 'CHEMBL1', 'molecule_structures': {
                    'standard_inchi_key': 'WRONG-KEY-N' if self.wrong_key else 'FULL-KEY-N'}}]}
            return {'activities': [], 'mechanisms': []}
        if db == 'pubmed':
            if command == 'search_pubmed':
                return ['10', '11']
            if command == 'fetch_article_abstracts':
                return [{'pmid': pmid, 'title': 'Synthetic article', 'abstract': 'Synthetic abstract'}
                        for pmid in args[1].split(',')]
        raise AssertionError((db, args))


class ValidationTests(unittest.TestCase):
    def test_verified_quote_and_offsets_with_page(self):
        row = validate_experiment(experiment(), SOURCE, 'paper')
        self.assertEqual(row['evidence_status'], 'text_supported')
        span = json.loads(row['evidence_checks'])['treatment_dose']['span']
        self.assertEqual(span['page'], 2)
        self.assertEqual(SOURCE[span['start']:span['end']], span['text'])

    def test_true_quote_cannot_support_fabricated_dose(self):
        row = validate_experiment(experiment(treatment_dose='99 uM'), SOURCE, 'paper')
        self.assertTrue(row['needs_review'])
        self.assertEqual(json.loads(row['evidence_checks'])['treatment_dose']['status'], 'needs_review')

    def test_hallucinated_quote_is_not_verified(self):
        row = experiment(field_evidence=json.dumps({'compound_name': 'Invented passage'}))
        self.assertTrue(validate_experiment(row, SOURCE, 'paper')['needs_review'])

    def test_numbered_model_evidence_and_named_combination(self):
        source = 'The combination of ABT199 and WEHI539 selectively killed senescent cells.'
        row = {'compound_name': 'ABT199 + WEHI539', 'experiment_model': 'senescent cells',
               'observed_result': 'selectively killed senescent cells',
               'field_evidence': json.dumps({'field_01': source, 'field_06': source, 'field_19': source})}
        checked = validate_experiment(row, source, 'paper')
        checks = json.loads(checked['evidence_checks'])
        self.assertEqual(checks['compound_name']['status'], 'text_supported')
        self.assertEqual(checks['observed_result']['status'], 'text_supported')

    def test_missing_dose_retains_negative_result(self):
        row = validate_experiment(experiment(treatment_dose=''), SOURCE, 'paper')
        self.assertEqual(row['observed_result'], 'No change was observed.')
        self.assertIn('treatment_dose', row['missing_fields'])

    def test_models_and_doses_have_distinct_experiment_ids(self):
        base = validate_experiment(experiment(), SOURCE, 'paper')['experiment_id']
        for changes in ({'treatment_dose': '4 uM'}, {'experiment_model': 'mice'}, {'treatment_duration': '72 h'}):
            self.assertNotEqual(base, validate_experiment(experiment(**changes), SOURCE, 'paper')['experiment_id'])

    def test_procedure_cannot_be_filled_from_common_knowledge(self):
        steps = [{'step': 1, 'action': 'Centrifuge at 1000 rpm', 'evidence_text': 'Centrifuge at 1000 rpm'}]
        row = validate_experiment(experiment(experimental_procedure=json.dumps(steps)), SOURCE, 'paper')
        self.assertTrue(row['needs_review'])
        self.assertIsNone(json.loads(row['procedure_steps'])[0]['evidence'])

    def test_whitespace_only_matching_and_no_invented_page(self):
        span = source_span('Dose was 2\n uM.', 'Dose was 2 uM.')
        self.assertIsNotNone(span)
        self.assertIsNone(span['page'])
        self.assertIsNone(source_span('Dose was 2 uM.', 'Dose was 3 uM.'))

    def test_untreated_control_and_background_are_not_tested_chemicals(self):
        control = validate_experiment(experiment(compound_name='no senolytics (untreated control)'), SOURCE, 'paper')
        background = validate_experiment(experiment(compound_name='Nutlin3a', evidence_origin='background'), SOURCE, 'paper')
        self.assertEqual(control['record_type'], 'nonchemical_intervention')
        self.assertEqual(background['record_type'], 'background_mention')



class EnrichmentTests(unittest.TestCase):
    def test_exact_structure_join_and_external_scope(self):
        skills = FakeSkills()
        result = enrich_compound('Quercetin', skills, source_pmid='10')
        self.assertEqual(result['status'], 'resolved')
        self.assertEqual(result['chembl']['id'], 'CHEMBL1')
        self.assertEqual([p['pmid'] for p in result['literature']['papers']], ['11'])
        self.assertIn('abstract', result['literature']['papers'][0]['evidence_scope'])
        self.assertTrue(any('molecule_structures__standard_inchi_key=FULL-KEY-N' in args for db,args in skills.calls))

    def test_ambiguous_identity_stops_downstream_enrichment(self):
        skills = FakeSkills(cids=[1, 2])
        self.assertEqual(enrich_compound('Q', skills)['status'], 'ambiguous')
        self.assertEqual(len(skills.calls), 1)

    def test_no_result_is_unresolved_not_guessed(self):
        self.assertEqual(enrich_compound('Q', FakeSkills(cids=[]))['status'], 'unresolved')

    def test_network_failure_is_error_not_no_result(self):
        self.assertEqual(enrich_compound('Q', FakeSkills(fail='pubchem'))['status'], 'error')

    def test_chembl_structure_conflict_does_not_merge(self):
        r = enrich_compound('Q', FakeSkills(wrong_key=True))
        self.assertEqual(r['chembl']['status'], 'unresolved')
        self.assertNotIn('id', r['chembl'])

    def test_paper_structure_conflict_is_visible(self):
        self.assertEqual(enrich_compound('Q', FakeSkills(), expected_inchikey='OTHER')['status'], 'identity_conflict')

    def test_local_codes_and_mixtures_require_identity_review(self):
        for name in ['Compound 1', 'Compound A', 'D+Q']:
            skills = FakeSkills()
            self.assertEqual(enrich_compound(name, skills)['status'], 'unresolved')
            self.assertEqual(skills.calls, [])

    def test_explicit_combination_component_facts_do_not_become_single_agent_results(self):
        r = enrich_compound('Quercetin + Dasatinib', FakeSkills())
        self.assertEqual(r['status'], 'combination_components_resolved')
        self.assertEqual(len(r['components']), 2)
        self.assertNotIn('cid', r['pubchem'])
        self.assertIn('do not attribute', r['interpretation'])

    def test_combination_alias_uses_only_a_grounded_definition_from_the_same_paper(self):
        support = json.dumps({'compound_name': {'status': 'text_supported'}})
        common = {'source_id': 'paper-a', 'evidence_checks': support, 'evidence_status': 'needs_review'}
        rows = [{**common, 'compound_name': 'Quercetin + Dasatinib (D+Q)'},
                {**common, 'compound_name': 'D+Q'},
                {**common, 'source_id': 'paper-b', 'compound_name': 'D+Q'}]
        with tempfile.TemporaryDirectory() as tmp:
            out = enrich_rows(rows, Path(tmp), skills=FakeSkills())
        self.assertEqual(out[1]['compound_name'], 'D+Q')
        self.assertEqual(out[1]['enrichment_status'], 'combination_components_resolved')
        self.assertEqual(out[2]['enrichment_status'], 'unresolved')

    def test_cell_transplant_intervention_is_retained_but_not_queried_as_a_chemical(self):
        row = validate_experiment(experiment(compound_name='senescent preadipocytes'), SOURCE, 'paper')
        self.assertEqual(row['record_type'], 'nonchemical_intervention')
        skills = FakeSkills()
        result = enrich_rows([row], Path('.'), skills=skills)[0]
        self.assertEqual(result['enrichment_status'], 'skipped_nonchemical_intervention')
        self.assertEqual(skills.calls, [])

    def test_background_mention_is_retained_without_database_lookup(self):
        row = validate_experiment(experiment(compound_name='Nutlin3a', evidence_origin='background'), SOURCE, 'paper')
        skills = FakeSkills()
        result = enrich_rows([row], Path('.'), skills=skills)[0]
        self.assertEqual(result['enrichment_status'], 'skipped_background_mention')
        self.assertEqual(skills.calls, [])

    def test_enrichment_never_overwrites_original_experiment(self):
        skills = FakeSkills()
        row = validate_experiment(experiment(), SOURCE, 'paper')
        with tempfile.TemporaryDirectory() as tmp:
            rows = enrich_rows([row, row], Path(tmp), skills=skills)
        self.assertEqual(rows[0]['observed_result'], row['observed_result'])
        self.assertEqual(len([1 for db,args in skills.calls if db=='pubchem' and args[0]=='resolve']), 1)
        self.assertEqual(len(rows), 2)

    def test_unverified_experiment_does_not_trigger_network(self):
        skills = FakeSkills()
        rows = enrich_rows([experiment()], Path('.'), skills=skills)
        self.assertEqual(rows[0]['enrichment_status'], 'skipped_unverified_experiment')
        self.assertEqual(skills.calls, [])

    def test_grounded_name_can_be_enriched_without_approving_uncertain_dose(self):
        row = validate_experiment(experiment(treatment_dose='99 uM'), SOURCE, 'paper')
        skills = FakeSkills()
        with tempfile.TemporaryDirectory() as tmp:
            result = enrich_rows([row], Path(tmp), skills=skills)[0]
        self.assertTrue(result['needs_review'])
        self.assertEqual(result['evidence_status'], 'needs_review')
        self.assertEqual(result['enrichment_status'], 'resolved')
        self.assertEqual(result['treatment_dose'], '99 uM')

    def test_partial_service_failure_preserves_successful_sources(self):
        r = enrich_compound('Quercetin', FakeSkills(fail='chembl'))
        self.assertEqual(r['status'], 'resolved')
        self.assertEqual(r['chembl']['status'], 'error')
        self.assertTrue(r['literature']['papers'])


class WorkflowTests(unittest.TestCase):
    def test_seed21_request_uses_documented_non_thinking_mode(self):
        from chem_pdf_extractor.llm import cloud_chat_completion
        response = io.BytesIO(json.dumps({'choices': [{'message': {'content': '{"ok":true}'}}]}).encode())
        with patch('urllib.request.urlopen', return_value=response) as request:
            self.assertEqual(cloud_chat_completion(
                'https://ark.cn-beijing.volces.com/api/v3', 'test-only',
                'doubao-seed-2-1-turbo-260628', [{'role': 'user', 'content': 'Return JSON'}], 30),
                '{"ok":true}')
        body = json.loads(request.call_args.args[0].data)
        self.assertEqual(body['reasoning_effort'], 'minimal')
        self.assertEqual(body['max_tokens'], 16000)

    def test_seed_evolving_request_has_enough_json_output_budget(self):
        from chem_pdf_extractor.llm import cloud_chat_completion
        response = io.BytesIO(json.dumps({'choices': [{'message': {'content': '{"ok":true}'}}]}).encode())
        with patch('urllib.request.urlopen', return_value=response) as request:
            self.assertEqual(cloud_chat_completion(
                'https://ark.cn-beijing.volces.com/api/v3', 'test-only',
                'doubao-seed-evolving', [{'role': 'user', 'content': 'Return JSON'}], 30),
                '{"ok":true}')
        body = json.loads(request.call_args.args[0].data)
        self.assertEqual(body['reasoning_effort'], 'minimal')
        self.assertEqual(body['max_tokens'], 16000)

    def test_evaluation_does_not_count_mixed_conditions_or_duplicate_hits(self):
        gold = [experiment(), experiment(treatment_dose='4 uM')]
        predicted = [experiment(), experiment(), experiment(treatment_dose='99 uM')]
        result = evaluate(predicted, gold)
        self.assertEqual(result['correct_records'], 1)
        self.assertEqual(result['experiment_precision'], 1/3)
        self.assertEqual(result['experiment_recall'], 1/2)

    def test_all_source_characters_reach_chunks(self):
        text = ''.join(str(i%10) for i in range(50001))
        chunks = list(text_chunks(text, 18000, 1200))
        restored = chunks[0] + ''.join(c[1200:] for c in chunks[1:])
        self.assertEqual(restored, text)

    def test_excel_does_not_silently_truncate_external_evidence(self):
        from chem_pdf_extractor.export import export_excel
        from openpyxl import load_workbook
        from types import SimpleNamespace
        import pandas as pd
        import warnings
        text = 'evidence ' * 8000 + '🧪' * 1000
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'evidence.xlsx'
            with warnings.catch_warnings():
                warnings.simplefilter('error', UserWarning)
                export_excel([{'external_evidence': text}], path, SimpleNamespace(pd=pd))
            book = load_workbook(path, read_only=True)
            headers, cells = list(book.active.iter_rows(values_only=True))
            parts = [value for key, value in zip(headers, cells) if str(key).startswith('external_evidence')]
            book.close()
        self.assertEqual(''.join(parts), text)
        self.assertTrue(all(len(part.encode('utf-16-le')) // 2 <= 32000 for part in parts))

    def test_combined_supplement_and_export_do_not_require_live_api(self):
        import pandas as pd
        from pydantic import Field, create_model
        runtime = RuntimeDeps(pd=pd, PdfReader=None, ChatPromptTemplate=None,
                              ChatOllama=None, Field=Field, create_model=create_model)
        paper = {'source': 'paper.txt', 'text': SOURCE, 'scope': 'local_document'}
        supplement = {'source': 'supplement.txt', 'text': 'Additional methods.', 'scope': 'local_document'}
        with tempfile.TemporaryDirectory() as tmp, patch(
            'chem_pdf_extractor.biomedical.extract_with_cloud_api', return_value=[experiment()]
        ), patch('urllib.request.urlopen', side_effect=AssertionError('No network expected')):
            out = Path(tmp)
            rows = extract_article(paper, [supplement], {'enrich_databases': False}, runtime, FakeSkills(), out)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['enrichment_status'], 'disabled')
            self.assertTrue((out/'experiments.xlsx').is_file())
            self.assertIn('Additional methods.', (out/'source.md').read_text())
            self.assertEqual(len(json.loads((out/'report.json').read_text())['records']), 1)
            with patch('chem_pdf_extractor.biomedical.extract_with_cloud_api', side_effect=AssertionError('must reuse checkpoint')):
                resumed = extract_article(paper, [supplement], {'enrich_databases': False}, runtime, FakeSkills(), out)
            self.assertEqual(len(resumed), 1)

    def test_cloud_chunks_run_concurrently_but_records_keep_source_order(self):
        import hashlib
        import threading
        import time
        import pandas as pd
        from pydantic import Field, create_model
        runtime = RuntimeDeps(pd=pd, PdfReader=None, ChatPromptTemplate=None,
                              ChatOllama=None, Field=Field, create_model=create_model)
        text = 'A' * 3000 + 'B' * 3000 + 'C' * 3000
        chunks = list(text_chunks('# Document: concurrent.txt\n' + text, 3000))
        def name(chunk):
            return 'molecule-' + hashlib.sha256(chunk.encode()).hexdigest()[:8]
        lock = threading.Lock()
        active = peak = 0
        def extract(_path, _config, _fields, _mapping, chunk):
            nonlocal active, peak
            with lock:
                active += 1
                peak = max(peak, active)
            time.sleep(0.01)
            with lock:
                active -= 1
            return [experiment(compound_name=name(chunk))]
        with tempfile.TemporaryDirectory() as tmp, patch(
            'chem_pdf_extractor.biomedical.extract_with_cloud_api', side_effect=extract
        ):
            rows = extract_article({'source': 'concurrent.txt', 'text': text, 'scope': 'local_document'}, [],
                                   {'enrich_databases': False, 'chunk_size': 3000, 'workers': 3},
                                   runtime, FakeSkills(), Path(tmp))
        self.assertGreaterEqual(peak, 2)
        self.assertEqual([row['compound_name'] for row in rows], [name(chunk) for chunk in chunks])

    def test_failed_large_chunks_split_and_successful_outputs_are_resumable(self):
        from chem_pdf_extractor.llm import CloudStructuredOutputError, TransientCloudAPIError
        import hashlib
        import pandas as pd
        from pydantic import Field, create_model
        runtime = RuntimeDeps(pd=pd, PdfReader=None, ChatPromptTemplate=None,
                              ChatOllama=None, Field=Field, create_model=create_model)
        paper = {'source': 'split.txt', 'text': 'A'*3000+'B'*2000, 'scope': 'local_document'}
        for error in [CloudStructuredOutputError]:
            with self.subTest(error=error.__name__), tempfile.TemporaryDirectory() as tmp:
                def extract(_path, _config, _fields, _mapping, chunk):
                    if len(chunk) > 2500:
                        raise error('synthetic failed large response')
                    return [experiment(compound_name='compound-'+hashlib.sha256(chunk.encode()).hexdigest()[:8])]
                config = {'enrich_databases': False, 'chunk_size': 4000, 'workers': 1}
                with patch('chem_pdf_extractor.biomedical.extract_with_cloud_api', side_effect=extract):
                    rows = extract_article(paper, [], config, runtime, FakeSkills(), Path(tmp))
                self.assertEqual(len(rows), 3)
                with patch('chem_pdf_extractor.biomedical.extract_with_cloud_api', side_effect=AssertionError('cached')):
                    resumed = extract_article(paper, [], config, runtime, FakeSkills(), Path(tmp))
                self.assertEqual(len(resumed), 3)

    def test_network_failure_does_not_recursively_split_the_paper(self):
        from chem_pdf_extractor.llm import TransientCloudAPIError
        import pandas as pd
        from pydantic import Field, create_model
        runtime=RuntimeDeps(pd=pd,PdfReader=None,ChatPromptTemplate=None,ChatOllama=None,Field=Field,create_model=create_model)
        events=[]
        config={'enrich_databases':False,'chunk_size':8000,'workers':1,
                'progress':lambda stage,**event:events.append((stage,event))}
        with tempfile.TemporaryDirectory() as tmp, patch('chem_pdf_extractor.biomedical.extract_with_cloud_api',
                side_effect=TransientCloudAPIError('timeout')) as call:
            with self.assertRaises(TransientCloudAPIError):
                extract_article({'source':'paper.txt','text':'a'*4000,'scope':'local_document'},[],config,runtime,FakeSkills(),Path(tmp))
            call.assert_called_once()
            self.assertEqual(events[0][1]['completed'],0)
            self.assertFalse((Path(tmp)/'report.json').exists())

    def test_doi_resolution_requires_unique_id(self):
        with self.assertRaisesRegex(ValueError, 'unique'):
            load_paper('DOI:10.1234/example', FakeSkills(), None)

    def test_query_is_bounded_and_has_aging_context(self):
        query = build_pubmed_query('Q"[all]', ['Q', 'A very long name'*20])
        self.assertIn('senescence[Title/Abstract]', query)
        self.assertNotIn('[all]', query)

    def test_skills_reject_unknown_database(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                ScienceSkills(Path(tmp)).call('unknown', 'query')

    def test_template_and_prompt_preserve_original_units_and_negative_results(self):
        labels = {f['label'] for f in ANTI_AGING_FIELDS}
        self.assertIn('statistical_evidence', labels)
        self.assertIn('experimental_procedure', labels)
        self.assertIn('阴性', BIOLOGY_RULES)
        self.assertIn('限定符', BIOLOGY_RULES)
        self.assertIn('claim_basis', labels)
        self.assertIn('不要为引言、讨论、综述和引用文献中的旧实验生成记录', BIOLOGY_RULES)


if __name__ == '__main__':
    unittest.main()
